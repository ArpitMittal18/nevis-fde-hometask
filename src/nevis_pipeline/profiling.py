from __future__ import annotations

import csv
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ENUM_COLUMNS = {
    "status", "advisor", "service rep", "segment", "risk profile",
    "fee schedule", "type", "attendee", "role", "office",
    "account type", "custodian",
}
TEXT_COLUMNS = {"name", "full name", "household", "client", "referred by", "account holder"}
DATE_COLUMNS = {"date", "client since", "as of date"}
CURRENCY_AMOUNT_COLUMNS = {"market value", "cost basis", "amount", "aum", "balance"}
CURRENCY_CODE_COLUMNS = {"currency", "currency code"}
IDENTIFIER_COLUMNS = {"advisor id", "account number"}
URL_COLUMNS = {"website", "url"}
MULTI_ENUM_COLUMNS = {"tags"}


def _normalized(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", name.lower()).strip()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_number(value: Any) -> bool:
    if isinstance(value, bool) or _is_blank(value):
        return False
    if isinstance(value, (int, float)):
        return True
    try:
        float(str(value).replace(",", ""))
        return True
    except ValueError:
        return False


def _is_date(value: Any) -> bool:
    if isinstance(value, (date, datetime)):
        return True
    text = str(value).strip()
    for fmt in ("%B %d, %Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            datetime.strptime(text, fmt)
            return True
        except ValueError:
            pass
    return False


def _string_values(values: list[Any]) -> list[str]:
    return sorted({str(value) for value in values if not _is_blank(value)})


def _classify(name: str, values: list[Any], row_count: int) -> str:
    normalized = _normalized(name)
    non_null = [value for value in values if not _is_blank(value)]
    if normalized in IDENTIFIER_COLUMNS or normalized.endswith(" id"):
        return "identifier"
    if normalized in DATE_COLUMNS or (non_null and all(_is_date(value) for value in non_null)):
        return "date"
    if normalized in CURRENCY_CODE_COLUMNS:
        return "currency_code"
    if normalized in CURRENCY_AMOUNT_COLUMNS or any(term in normalized for term in ("market value", "cost basis", "amount", "balance", "aum")):
        return "currency_amount"
    if normalized in URL_COLUMNS:
        return "url"
    if normalized in MULTI_ENUM_COLUMNS:
        return "multi_enum"
    if normalized in ENUM_COLUMNS:
        return "enum"
    if non_null and all(isinstance(value, bool) for value in non_null):
        return "boolean"
    if non_null and all(_is_number(value) for value in non_null):
        return "number"
    if normalized in TEXT_COLUMNS:
        return "text"
    distinct = len(_string_values(non_null))
    if distinct and distinct <= 20 and distinct / max(row_count, 1) <= 0.25:
        return "enum"
    return "text"


def _profile_columns(headers: list[str], rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    profiles = []
    for name in headers:
        raw_values = [row.get(name) for row in rows]
        non_null = [value for value in raw_values if not _is_blank(value)]
        column_type = _classify(name, raw_values, len(rows))
        profile: dict[str, Any] = {
            "name": name,
            "type": column_type,
            "row_count": len(rows),
            "non_null_count": len(non_null),
            "null_count": len(rows) - len(non_null),
            "distinct_count": len(_string_values(non_null)),
        }
        if column_type == "multi_enum":
            profile["values"] = sorted({part.strip() for value in non_null for part in str(value).split(",") if part.strip()})
        elif column_type in {"enum", "currency_code", "boolean"}:
            profile["values"] = _string_values(non_null)
        else:
            profile["examples"] = _string_values(non_null)[:5]
        profiles.append(profile)
    return profiles


def _dataset_name(path: Path) -> str:
    name = path.name.lower()
    if name.startswith("clients "):
        return "clients"
    if name.startswith("meetings "):
        return "meetings"
    if name == "advisor_roster.csv":
        return "advisor_roster"
    if name == "custodian_positions.xlsx":
        return "custodian_positions"
    return re.sub(r"[^a-z0-9]+", "_", path.stem.lower()).strip("_")


def profile_sources(sources_dir: Path) -> dict[str, Any]:
    sources_dir = Path(sources_dir)
    csv_files = []
    for path in sorted(sources_dir.rglob("*.csv")):
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            rows = list(reader)
            headers = list(reader.fieldnames or [])
        csv_files.append({
            "dataset": _dataset_name(path),
            "path": str(path.relative_to(sources_dir.parent)),
            "row_count": len(rows),
            "columns": _profile_columns(headers, rows),
        })

    workbooks = []
    for path in sorted(sources_dir.glob("*.xlsx")):
        workbook = load_workbook(path, data_only=True, read_only=True)
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value) for value in next(iterator)]
            rows = [dict(zip(headers, values)) for values in iterator]
            workbooks.append({
                "dataset": _dataset_name(path),
                "path": str(path.relative_to(sources_dir.parent)),
                "sheet": sheet.title,
                "row_count": len(rows),
                "columns": _profile_columns(headers, rows),
            })

    return {
        "description": "Source schema and value profile generated before canonical mapping",
        "type_legend": {
            "identifier": "Stable source identifier, not a category",
            "enum": "Single categorical value; values lists the observed domain",
            "multi_enum": "Comma-separated categorical values",
            "date": "Date or date-like source value",
            "currency_code": "ISO-like currency category; values lists observed currencies",
            "currency_amount": "Monetary amount requiring an associated currency",
            "number": "Non-currency numeric value",
            "boolean": "True/false value",
            "url": "Web address",
            "text": "Free text, name, or reference-like value",
        },
        "csv_files": csv_files,
        "workbooks": workbooks,
    }


def render_profile_markdown(profile: dict[str, Any]) -> str:
    lines = ["# Source Data Profile", "", "Generated before canonical mapping. Types describe the source columns, not the final Nevis schema.", ""]
    datasets = profile["csv_files"] + profile["workbooks"]
    for dataset in datasets:
        suffix = f" — sheet `{dataset['sheet']}`" if dataset.get("sheet") else ""
        lines.extend([f"## {dataset['dataset']}{suffix}", "", f"Source: `{dataset['path']}` · Rows: **{dataset['row_count']}**", "", "| Column | Type | Nulls | Distinct | Observed values / examples |", "|---|---:|---:|---:|---|"])
        for column in dataset["columns"]:
            displayed = column.get("values", column.get("examples", []))
            values = ", ".join(f"`{value}`" for value in displayed)
            lines.append(f"| `{column['name']}` | `{column['type']}` | {column['null_count']} | {column['distinct_count']} | {values} |")
        lines.append("")
    lines.extend(["## Type legend", ""])
    for name, description in profile["type_legend"].items():
        lines.append(f"- `{name}` — {description}")
    lines.append("")
    return "\n".join(lines)
