from __future__ import annotations

import csv
import hashlib
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import SourceRecord, SourceRef, StagingBundle


CLIENT_COLUMNS = {"Name", "Status", "Advisor", "Service Rep", "Household"}
MEETING_COLUMNS = {"Name", "Client", "Type", "Date", "Attendee"}
ACCOUNT_COLUMNS = {"Account_Number", "Account_Holder", "Account_Type", "Market_Value", "Currency", "As_Of_Date", "Custodian"}
ADVISOR_COLUMNS = {"advisor_id", "full_name", "role", "office"}


def _checksum(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_csv(path: Path, required: set[str]) -> list[SourceRecord]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"{path}: missing required columns {sorted(missing)}")
        return [SourceRecord(dict(row), SourceRef(str(path), row=index)) for index, row in enumerate(reader, start=2)]


def _page_index(folder: Path) -> dict[str, tuple[str, str]]:
    result: dict[str, tuple[str, str]] = {}
    for path in folder.glob("*.md"):
        text = path.read_text(encoding="utf-8")
        lines = text.splitlines()
        if not lines:
            continue
        title = lines[0].removeprefix("# ").strip()
        properties = ("Status:", "Advisor:", "Service Rep:", "Household:", "Segment:", "Client Since:", "Website:", "Risk Profile:", "Fee Schedule:", "Referred By:", "Tags:", "Client:", "Type:", "Date:", "Attendee:")
        body = "\n".join(line for line in lines[1:] if line.strip() and not line.startswith(properties)).strip()
        result[title] = (str(path), body)
    return result


def _attach_pages(records: Iterable[SourceRecord], pages: dict[str, tuple[str, str]]) -> None:
    for record in records:
        page = pages.get(record.values["Name"])
        if page:
            record.page_path, record.page_body = page


def ingest_sources(sources_dir: Path) -> StagingBundle:
    sources_dir = Path(sources_dir)
    notion = sources_dir / "notion_export"
    clients_path = next(notion.glob("Clients *.csv"))
    meetings_path = next(notion.glob("Meetings *.csv"))
    client_folder = next(path for path in notion.glob("Clients *") if path.is_dir())
    meeting_folder = next(path for path in notion.glob("Meetings *") if path.is_dir())

    clients = _read_csv(clients_path, CLIENT_COLUMNS)
    meetings = _read_csv(meetings_path, MEETING_COLUMNS)
    advisors_path = sources_dir / "advisor_roster.csv"
    advisors = _read_csv(advisors_path, ADVISOR_COLUMNS)
    _attach_pages(clients, _page_index(client_folder))
    _attach_pages(meetings, _page_index(meeting_folder))

    workbook_path = sources_dir / "custodian_positions.xlsx"
    worksheet = load_workbook(workbook_path, data_only=True, read_only=True)["Positions"]
    headers = [cell.value for cell in next(worksheet.iter_rows())]
    missing = ACCOUNT_COLUMNS - set(headers)
    if missing:
        raise ValueError(f"{workbook_path}: missing required columns {sorted(missing)}")
    accounts = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
        values = {key: cell.value for key, cell in zip(headers, row)}
        accounts.append(SourceRecord(values, SourceRef(str(workbook_path), row=row_number)))

    slack_path = sources_dir / "ops_slack_thread.md"
    inputs = [clients_path, meetings_path, advisors_path, workbook_path, slack_path]
    inputs.extend(Path(record.page_path) for record in clients + meetings if record.page_path)
    return StagingBundle(
        clients=clients,
        meetings=meetings,
        accounts=accounts,
        advisors=advisors,
        slack_text=slack_path.read_text(encoding="utf-8"),
        input_checksums={str(path.relative_to(sources_dir.parent)): _checksum(path) for path in sorted(set(inputs))},
    )
